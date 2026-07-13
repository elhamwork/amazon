from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import config

COLUMNS = [
    ("amazon_title", "Product Title"),
    ("aliexpress_price", "AliExpress Price"),
    ("amazon_price", "Amazon Price"),
    ("fba_fee", "FBA Fee"),
    ("net_profit", "Net Profit"),
    ("margin_pct", "Margin %"),
    ("bsr", "BSR"),
    ("review_count", "Review Count"),
    ("rating", "Rating"),
    ("seller_count", "Seller Count"),
    ("fba_listed", "FBA Listed"),
    ("keepa_price_trend", "Price Trend"),
    ("keepa_bsr_trend", "BSR Trend"),
    ("selleramp_margin_pct", "SAS Margin %"),
    ("selleramp_roi_pct", "SAS ROI %"),
    ("selleramp_monthly_sales_est", "SAS Monthly Sales"),
    ("helium10_monthly_sales_est", "H10 Monthly Sales"),
    ("helium10_monthly_revenue_est", "H10 Monthly Revenue"),
    ("helium10_review_velocity", "H10 Review Velocity"),
    ("datadive_top_keyword", "Top Keyword"),
    ("datadive_search_volume", "Keyword Search Volume"),
    ("datadive_ppc_bid", "Keyword PPC Bid"),
    ("score", "Score"),
]

SCORE_FILLS = {
    "GO": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "WATCH": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "SKIP": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
}


def write_excel(rows: list[dict], output_path: str = None) -> str:
    output_path = output_path or config.OUTPUT_XLSX
    wb = Workbook()
    ws = wb.active
    ws.title = "FBA Research"

    header_font = Font(bold=True)
    for col_idx, (_, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, start=2):
        score = row.get("score", "SKIP")
        fill = SCORE_FILLS.get(score)
        for col_idx, (key, _) in enumerate(COLUMNS, start=1):
            value = row.get(key)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if fill:
                cell.fill = fill

    for col_idx, (key, label) in enumerate(COLUMNS, start=1):
        max_len = len(label)
        for row in rows:
            value = row.get(key)
            if value is not None:
                max_len = max(max_len, len(str(value)))
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = max_len + 4

    ws.freeze_panes = "A2"
    wb.save(output_path)
    return output_path
